from __future__ import annotations

import re
from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


HEADERS = (
    "Order",
    "Date",
    "Customer",
    "Store / Channel",
    "Total",
    "Payment",
    "Order Status",
    "Items",
    "Delivery / Carrier",
    "Tracking No",
)


def _pick(data: dict[str, Any], *keys: str, default: Any = None) -> Any:
    if not isinstance(data, dict):
        return default
    lowered = {str(key).lower(): value for key, value in data.items()}
    for key in keys:
        value = lowered.get(key.lower())
        if value not in (None, ""):
            return value
    return default


def _shown(value: Any) -> bool:
    return value not in (None, "", [], {})


def _payload(row: dict[str, Any]) -> dict[str, Any]:
    value = row.get("details", row)
    if not isinstance(value, dict):
        return {}
    result = value.get("result", value)
    return result if isinstance(result, dict) else {}


def _order_reference(payload: dict[str, Any], fallback: int) -> str:
    order = payload.get("order") or payload.get("Order") or payload
    return str(
        _pick(
            order,
            "orderNo",
            "refNo",
            "orderId",
            default=f"Order {fallback}",
        )
    )


def _excel_row(row: dict[str, Any], group_name: str, number: int) -> list[Any]:
    payload = _payload(row)
    summary = row.get("summary") if isinstance(row, dict) else {}
    summary = summary if isinstance(summary, dict) else {}
    order = payload.get("order") or payload.get("Order") or payload
    order = order if isinstance(order, dict) else {}
    address = payload.get("orderAddress") or payload.get("OrderAddress") or {}
    address = address if isinstance(address, dict) else {}
    items = payload.get("orderItems") or payload.get("OrderItems") or []
    items = items if isinstance(items, list) else []

    payment_id = _pick(order, "paymentStatusId")
    payment = _pick(
        summary,
        "paymentStatus",
        default=_pick(order, "paymentStatus", "paymentStatusName"),
    )
    if not _shown(payment):
        payment = {1: "Unpaid", 2: "Paid"}.get(payment_id, "-")

    delivery_status = _pick(
        summary,
        "carrierTrackingStatus",
        "deliveryStatus",
        default=_pick(order, "carrierTrackingStatus"),
    )
    carrier = _pick(
        summary,
        "carrierName",
        "deliveryMethod",
        default=_pick(order, "carrierName"),
    )
    delivery = " / ".join(
        str(value) for value in (delivery_status, carrier) if _shown(value)
    ) or "-"
    source = _pick(summary, "storeName", "saleChannelName", "channelName")

    return [
        _order_reference(payload, number),
        _pick(order, "orderDate", default=_pick(summary, "orderDate")) or "-",
        _pick(address, "customerName", default=_pick(summary, "customerName")) or "-",
        source or group_name,
        _pick(order, "amount", default=_pick(summary, "amount")) or 0,
        payment,
        _pick(
            summary,
            "orderStatus",
            "status",
            default=_pick(order, "orderStatus", "status"),
        ) or delivery_status or "-",
        _pick(order, "itemsCount", default=len(items) if items else "-"),
        delivery,
        _pick(
            summary,
            "carrierTrackingNo",
            "trackingNo",
            default=_pick(order, "carrierTrackingNo", "trackingNo"),
        ) or "-",
    ]


def safe_excel_filename(title: str) -> str:
    clean = re.sub(r"[^A-Za-z0-9_-]+", "_", title).strip("_")
    return f"{clean[:80] or 'shipra_orders'}.xlsx"


def build_orders_excel(
    groups: list[dict[str, Any]],
    *,
    title: str,
    filter_text: str,
) -> bytes:
    """Create a professional workbook containing the verified order result."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Orders"

    total_orders = sum(len(group.get("rows") or []) for group in groups)
    sheet.append([title])
    sheet.append([filter_text])
    sheet.append([
        f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')} | "
        f"Total orders: {total_orders}"
    ])
    sheet.append([])
    sheet.append(list(HEADERS))

    number = 0
    for group in groups:
        group_name = str(group.get("name") or "All stores")
        for row in group.get("rows") or []:
            number += 1
            sheet.append(_excel_row(row, group_name, number))

    last_column = get_column_letter(len(HEADERS))
    sheet.merge_cells(f"A1:{last_column}1")
    sheet.merge_cells(f"A2:{last_column}2")
    sheet.merge_cells(f"A3:{last_column}3")
    sheet["A1"].font = Font(size=16, bold=True, color="22242A")
    sheet["A2"].font = Font(size=10, color="696D78")
    sheet["A3"].font = Font(size=9, color="696D78")

    header_fill = PatternFill("solid", fgColor="563AD5")
    for cell in sheet[5]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sheet.freeze_panes = "A6"
    sheet.auto_filter.ref = f"A5:{last_column}{max(sheet.max_row, 5)}"
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[5].height = 24

    widths = (20, 22, 24, 24, 14, 15, 22, 10, 28, 22)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=6):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet["E"][5:]:
        cell.number_format = "#,##0.00"

    if total_orders:
        table = Table(
            displayName="ShipraOrders",
            ref=f"A5:{last_column}{sheet.max_row}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium4",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)

    metadata = workbook.create_sheet("Report Info")
    metadata.append(["Field", "Value"])
    metadata.append(["Title", title])
    metadata.append(["Filters", filter_text])
    metadata.append(["Total Orders", total_orders])
    metadata.append(["Generated", datetime.now().strftime("%Y-%m-%d %H:%M")])
    metadata.freeze_panes = "A2"
    metadata.column_dimensions["A"].width = 20
    metadata.column_dimensions["B"].width = 70
    for cell in metadata[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
